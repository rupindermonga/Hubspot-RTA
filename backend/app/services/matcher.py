"""Address matching engine — Hubspot vs RTA.

Pure-Python port of the Streamlit prototype's matching pipeline. Takes pandas
DataFrames in, returns a result dict + a color-coded XLSX byte stream out.

Matching pipeline (passes from safest to riskiest):
  1. Exact          — street + postal code match after normalization
  2. Direction strip — same as 1 ignoring trailing N/S/E/W
  3. Fuzzy/alias    — known street-name variants treated as one
  4. Unit suffix    — "-U1", "-U2" stripped from address numbers
  5. Street-only    — opt-in; risky because different towns can share a street name
"""
from __future__ import annotations

import io
import re
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill


# ── Constants ────────────────────────────────────────────────────────────────

ABBREVS = {
    'STREET': 'ST', 'ROAD': 'RD', 'DRIVE': 'DR', 'AVENUE': 'AVE',
    'BOULEVARD': 'BLVD', 'CRESCENT': 'CRES', 'COURT': 'CRT', 'PLACE': 'PL',
    'LANE': 'LN', 'CIRCLE': 'CIR', 'TERRACE': 'TERR', 'HIGHWAY': 'HWY',
    'TRAIL': 'TRL', 'SQUARE': 'SQ', 'PARKWAY': 'PKY', 'WAY': 'WAY',
    'CLOSE': 'CL', 'GROVE': 'GRV', 'HEIGHTS': 'HTS', 'RIDGE': 'RDG',
    'NORTH': 'N', 'SOUTH': 'S', 'EAST': 'E', 'WEST': 'W',
    'REG': 'REGIONAL',
}

# Default street-name aliases (variant → canonical). Extend as new variants
# surface in real data. If editing per-user becomes a need, move to a DB table
# and a route at /api/aliases — for now this is shared across all users.
DEFAULT_ALIASES: list[tuple[str, str]] = [
    ('PANACHE N SHR RD', 'PANACHE NSHORE RD'),
    ('PANACHE N SHORE RD', 'PANACHE NSHORE RD'),
    ('PANACHE NORTHSHORE RD', 'PANACHE NSHORE RD'),
    ('A PANACHE N SHORE RD', 'PANACHE NSHORE RD'),
    ('A PANACHE N SHR RD', 'PANACHE NSHORE RD'),
    ('PANACHE SHORE RD', 'PANACHE NSHORE RD'),
    ('NORTHSHORE RD', 'PANACHE NSHORE RD'),
    ('N SHORE RD', 'PANACHE NSHORE RD'),
    ('PENACHE NORTHSHORE RD', 'PANACHE NSHORE RD'),
    ('PENACHE N SHORE RD', 'PANACHE NSHORE RD'),
    ('HENNESSY RD', 'HENNESSEY RD'),
    ('OLD SYLVAIN VALLEY HILL RD', 'OLD SLYVAN VALLEY HILL RD'),
    ('LITTLE PENAGE LAKE RD', 'LITTLE PANACHE RD'),
    ('REGIONAL 10 RD', 'REGIONAL RD 10'),
    ('FINDLAY HILL RD', 'FINDLAY RD'),
    ('FINDLAY HILL RD E', 'FINDLAY RD E'),
    ('FINDLAY HILL RD W', 'FINDLAY RD W'),
]

FORMULA_PREFIXES = ('=', '+', '-', '@', '\t', '\r')


# ── Sanitization (Excel formula injection guard) ─────────────────────────────

def sanitize_cell(val):
    if isinstance(val, str) and val and val[0] in FORMULA_PREFIXES:
        return "'" + val
    return val


def sanitize_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    df_clean = df.copy()
    for col in df_clean.select_dtypes(include='object').columns:
        df_clean[col] = df_clean[col].apply(lambda v: sanitize_cell(v) if isinstance(v, str) else v)
    return df_clean


# ── Normalization helpers ────────────────────────────────────────────────────

def clean_address(s) -> str:
    """Strip PO Box, RR, Suite, Apt, Unit noise."""
    s = str(s).strip()
    s = re.sub(r'^(?:P[\s.]?O[\s.]?\s*)?BOX\s*#?\s*\d*[\s,/]*', '', s, flags=re.IGNORECASE)
    s = re.sub(r'^SUITE\s+\d+\s*', '', s, flags=re.IGNORECASE)
    s = re.sub(r'^RR\s*#?\s*\d+[\s,]*', '', s, flags=re.IGNORECASE)
    s = re.sub(r'[\s,]+(?:P[\s.]?O[\s.]?\s*)?BOX\s*#?\s*\d*.*$', '', s, flags=re.IGNORECASE)
    s = re.sub(r'[\s,]+RR\s*#?\s*\d+.*$', '', s, flags=re.IGNORECASE)
    s = re.sub(r'[\s,]+(?:SUITE|APT|UNIT)\s*#?\s*\w*.*$', '', s, flags=re.IGNORECASE)
    return s.strip(' ,/')


def normalize(s) -> str:
    s = clean_address(s)
    s = s.upper().replace('.', '')
    s = re.sub(r'[^A-Z0-9 ]', '', s)
    s = re.sub(r' +', ' ', s).strip()
    words = s.split()
    words = [ABBREVS.get(w, w) for w in words]
    return ' '.join(words)


def strip_direction(s: str) -> str:
    """Only strip trailing direction if street name has 3+ words.
    Prevents '123 N ST' from losing the 'N' which IS the street name."""
    words = s.split()
    if len(words) >= 3 and words[-1] in ('N', 'S', 'E', 'W'):
        return ' '.join(words[:-1])
    return s


def strip_unit(s: str) -> str:
    """Strip unit suffixes like -U1, U2 from address numbers.
    '97U2 PIONEER RD' -> '97 PIONEER RD'"""
    return re.sub(r'^(\d+)[- ]?U\d+', r'\1', s)


def norm_pc(s) -> str:
    """Normalize Canadian postal code prefix (first 3 chars). Fixes O↔0 / I↔1 typos."""
    s = str(s).strip().upper().replace(' ', '')
    if s in ('NAN', 'NONE', 'N/A', 'NA', 'NULL', ''):
        return ''
    corrected = []
    for i, c in enumerate(s[:6]):
        if i in (1, 3, 5):
            if c == 'O': c = '0'
            elif c == 'I': c = '1'
        elif i in (0, 2, 4):
            if c == '0': c = 'O'
        corrected.append(c)
    s = ''.join(corrected)
    return s[:3]


def apply_canonical(street_full: str, canonical_map: dict[str, str]) -> str:
    m = re.match(r'^(\d+[A-Z]?\s+)(.*)', street_full)
    if m:
        hnum, sname = m.group(1), m.group(2)
    else:
        hnum, sname = '', street_full
    sname = canonical_map.get(sname, sname)
    return hnum + sname


# ── Run the matching pipeline ────────────────────────────────────────────────

def run_match(
    df_hub: pd.DataFrame,
    df_rta: pd.DataFrame,
    *,
    hub_street_col: str,
    hub_pc_col: str,
    rta_addr_no_col: str,
    rta_street_col: str,
    rta_locality_col: str,
    rta_pc_col: str,
    rta_status_col: str,
    enable_no_pc: bool = False,
    aliases: Iterable[tuple[str, str]] | None = None,
) -> dict:
    """Match Hubspot rows against the RTA database.

    Returns a dict with:
      - stats (dict)             counts for the dashboard
      - conflicts (list of dict) RTA addresses that have conflicting statuses
      - flagged (list of dict)   non-exact matches needing visual review
      - rta_not_in_hubspot_preview (list of dict)
      - excel_bytes (bytes)      color-coded XLSX (Hubspot + RTA sheets)
    """
    canonical_map = dict(aliases) if aliases is not None else dict(DEFAULT_ALIASES)

    df_hub = df_hub.copy()
    df_rta = df_rta.copy()

    # Build RTA combined full address: "AddressNo StreetName Locality PostalCode"
    df_rta['_rta_full'] = (
        df_rta[rta_addr_no_col].fillna('').astype(str).str.strip() + ' ' +
        df_rta[rta_street_col].fillna('').astype(str).str.strip() + ' ' +
        df_rta[rta_locality_col].fillna('').astype(str).str.strip() + ' ' +
        df_rta[rta_pc_col].fillna('').astype(str).str.strip()
    ).str.strip()

    # Normalize Hubspot
    df_hub['_street'] = df_hub[hub_street_col].fillna('').apply(normalize)
    df_hub['_street_canon'] = df_hub['_street'].apply(lambda s: apply_canonical(s, canonical_map))
    df_hub['_pc'] = df_hub[hub_pc_col].fillna('').apply(norm_pc)

    # Normalize RTA: combine AddressNo + StreetName for matching key
    df_rta['_street'] = (
        df_rta[rta_addr_no_col].fillna('').astype(str) + ' ' +
        df_rta[rta_street_col].fillna('')
    ).apply(normalize)
    df_rta['_street_canon'] = df_rta['_street'].apply(lambda s: apply_canonical(s, canonical_map))
    df_rta['_pc'] = df_rta[rta_pc_col].fillna('').apply(norm_pc)

    # Build key variants
    for df in (df_hub, df_rta):
        df['_k_exact'] = df['_street'] + '|' + df['_pc']
        df['_k_dir'] = df['_street'].apply(strip_direction) + '|' + df['_pc']
        df['_k_canon'] = df['_street_canon'] + '|' + df['_pc']
        df['_k_canon_dir'] = df['_street_canon'].apply(strip_direction) + '|' + df['_pc']
        df['_k_unit'] = df['_street'].apply(strip_unit) + '|' + df['_pc']
        df['_k_unit_dir'] = df['_street'].apply(strip_unit).apply(strip_direction) + '|' + df['_pc']

    # Detect duplicate keys with conflicting statuses
    dup_check = df_rta.groupby('_k_exact')[rta_status_col].nunique()
    conflict_keys = dup_check[dup_check > 1]
    conflict_key_set = set(conflict_keys.index)

    conflict_detail = []
    for key in list(conflict_keys.index)[:50]:
        rows = df_rta[df_rta['_k_exact'] == key][[rta_addr_no_col, rta_street_col, rta_pc_col, rta_status_col]]
        for _, r in rows.iterrows():
            conflict_detail.append({
                'address': f"{r[rta_addr_no_col]} {r[rta_street_col]}",
                'postal_code': str(r[rta_pc_col]),
                'status': str(r[rta_status_col]),
                'key': key,
            })

    # Build lookups, overriding conflicting keys with concatenated candidates
    conflict_addr_map: dict[str, str] = {}
    conflict_status_map: dict[str, str] = {}
    for key in conflict_key_set:
        rows = df_rta[df_rta['_k_exact'] == key]
        addrs = rows['_rta_full'].dropna().unique()
        statuses = rows[rta_status_col].dropna().unique()
        conflict_addr_map[key] = ' | '.join(str(a) for a in addrs)
        conflict_status_map[key] = ' | '.join(str(s) for s in statuses)

    lookup_addr: dict[str, pd.Series] = {}
    lookup_status: dict[str, pd.Series] = {}
    for key_col in ['_k_exact', '_k_dir', '_k_canon', '_k_canon_dir', '_k_unit', '_k_unit_dir']:
        deduped = df_rta.drop_duplicates(subset=key_col).set_index(key_col)
        addr_series = deduped['_rta_full'].copy()
        status_series = deduped[rta_status_col].fillna('').astype(str).copy()
        for ck in conflict_key_set:
            if ck in addr_series.index:
                addr_series[ck] = f"CONFLICT: {conflict_addr_map[ck]}"
            if ck in status_series.index:
                status_series[ck] = f"CONFLICT: {conflict_status_map[ck]}"
        lookup_addr[key_col] = addr_series
        lookup_status[key_col] = status_series

    # Initialize output columns
    df_hub['RTA Address'] = pd.Series(dtype='object')
    df_hub['RTA Status'] = pd.Series(dtype='object')
    df_hub['_match_type'] = ''

    passes = [
        ('_k_exact', 'exact'),
        ('_k_dir', 'direction_strip'),
        ('_k_canon', 'fuzzy'),
        ('_k_canon_dir', 'fuzzy'),
        ('_k_unit', 'fuzzy'),
        ('_k_unit_dir', 'fuzzy'),
    ]

    for key_col, mtype in passes:
        unmatched = df_hub['RTA Address'].isna()
        mapped_addr = df_hub.loc[unmatched, key_col].map(lookup_addr[key_col])
        mapped_status = df_hub.loc[unmatched, key_col].map(lookup_status[key_col])
        matched_mask = mapped_addr.notna()
        if matched_mask.any():
            df_hub.loc[mapped_addr[matched_mask].index, 'RTA Address'] = mapped_addr[matched_mask].values
            df_hub.loc[mapped_status[matched_mask].index, 'RTA Status'] = mapped_status[matched_mask].values
            newly_matched = unmatched & df_hub['RTA Address'].notna() & (df_hub['_match_type'] == '')
            df_hub.loc[newly_matched, '_match_type'] = mtype

    # Mark conflict-matched rows
    for idx in df_hub[df_hub['RTA Address'].notna()].index:
        key = df_hub.loc[idx, '_k_exact']
        if key in conflict_key_set and df_hub.loc[idx, '_match_type'] == 'exact':
            df_hub.loc[idx, '_match_type'] = 'conflict'

    # Pass 5: street-only (no postal code) → RED (opt-in)
    if enable_no_pc:
        r_lookup_addr: dict[str, str] = {}
        r_lookup_status: dict[str, str] = {}
        r_lookup_addr_dir: dict[str, str] = {}
        r_lookup_status_dir: dict[str, str] = {}
        for i in range(len(df_rta)):
            addr_val = df_rta.iloc[i]['_rta_full']
            status_val = str(df_rta.iloc[i].get(rta_status_col, ''))
            for st_key in (df_rta.iloc[i]['_street'], df_rta.iloc[i]['_street_canon']):
                if st_key and st_key not in r_lookup_addr:
                    r_lookup_addr[st_key] = addr_val
                    r_lookup_status[st_key] = status_val
            for st_key in (strip_direction(df_rta.iloc[i]['_street']), strip_direction(df_rta.iloc[i]['_street_canon'])):
                if st_key and st_key not in r_lookup_addr_dir:
                    r_lookup_addr_dir[st_key] = addr_val
                    r_lookup_status_dir[st_key] = status_val

        unmatched = df_hub['RTA Address'].isna()
        for idx in df_hub[unmatched].index:
            h_st = df_hub.loc[idx, '_street']
            h_st_canon = df_hub.loc[idx, '_street_canon']
            for la, ls, key in [
                (r_lookup_addr, r_lookup_status, h_st),
                (r_lookup_addr, r_lookup_status, h_st_canon),
                (r_lookup_addr_dir, r_lookup_status_dir, strip_direction(h_st)),
                (r_lookup_addr_dir, r_lookup_status_dir, strip_direction(h_st_canon)),
            ]:
                if key in la:
                    df_hub.loc[idx, 'RTA Address'] = la[key]
                    df_hub.loc[idx, 'RTA Status'] = ls.get(key, '')
                    df_hub.loc[idx, '_match_type'] = 'no_pc'
                    break

    # Reverse lookup — find RTA addresses NOT in Hubspot
    matched_hub_keys: set[str] = set()
    key_cols_list = ['_k_exact', '_k_dir', '_k_canon', '_k_canon_dir', '_k_unit', '_k_unit_dir']
    matched_rows = df_hub[df_hub['RTA Address'].notna()]
    for key_col in key_cols_list:
        matched_hub_keys.update(matched_rows[key_col].dropna().unique())

    def _rta_in_hub(row) -> str:
        for kc in key_cols_list:
            if row[kc] in matched_hub_keys:
                return 'Yes'
        return 'No'

    df_rta['In Hubspot'] = df_rta.apply(_rta_in_hub, axis=1)
    rta_in_hub = int((df_rta['In Hubspot'] == 'Yes').sum())
    rta_not_in_hub = int((df_rta['In Hubspot'] == 'No').sum())

    # Stats
    exact_count = int((df_hub['_match_type'] == 'exact').sum())
    yellow_count = int(df_hub['_match_type'].isin(['fuzzy', 'direction_strip']).sum())
    orange_count = int((df_hub['_match_type'] == 'conflict').sum())
    red_count = int((df_hub['_match_type'] == 'no_pc').sum())
    hub_matched = int(df_hub['RTA Address'].notna().sum())
    hub_unmatched = len(df_hub) - hub_matched

    # Flagged matches for review
    flagged_df = df_hub[df_hub['_match_type'].isin(['fuzzy', 'direction_strip', 'no_pc', 'conflict'])][
        [hub_street_col, hub_pc_col, 'RTA Address', 'RTA Status', '_match_type']
    ].copy()
    flagged = [
        {
            'street_address': str(r[hub_street_col]),
            'postal_code': str(r[hub_pc_col]),
            'rta_address': str(r['RTA Address']),
            'rta_status': str(r['RTA Status']),
            'match_type': str(r['_match_type']),
        }
        for _, r in flagged_df.iterrows()
    ]

    # RTA-not-in-Hubspot preview (first 50)
    rta_not_matched_df = df_rta[df_rta['In Hubspot'] == 'No'][
        [rta_addr_no_col, rta_street_col, rta_locality_col, rta_pc_col, rta_status_col]
    ].head(50)
    rta_not_in_hubspot_preview = rta_not_matched_df.to_dict(orient='records')

    # Hubspot output preview: first 20 rows showing source columns + appended RTA fields
    hub_preview_df = df_hub[[hub_street_col, hub_pc_col, 'RTA Address', 'RTA Status']].head(20).copy()
    # Coerce NaN → '' so the JSON serializer doesn't drop or stringify them weirdly
    hub_preview_df = hub_preview_df.fillna('')
    hub_output_preview = hub_preview_df.to_dict(orient='records')

    # ── Build the colored XLSX ────────────────────────────────────────────
    match_type_series = df_hub['_match_type'].copy()
    df_hub_out = df_hub.drop(columns=[c for c in df_hub.columns if c.startswith('_')])
    df_rta_out = df_rta.drop(columns=[c for c in df_rta.columns if c.startswith('_')])

    df_hub_out = sanitize_dataframe(df_hub_out)
    df_rta_out = sanitize_dataframe(df_rta_out)

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
        df_hub_out.to_excel(writer, sheet_name='Hubspot', index=False)
        df_rta_out.to_excel(writer, sheet_name='RTA', index=False)
    buffer.seek(0)

    wb = load_workbook(buffer)

    # Color Hubspot RTA Address / RTA Status cells per match type
    ws_hub = wb['Hubspot']
    rta_addr_col_idx = None
    rta_status_col_idx = None
    for cell in ws_hub[1]:
        if cell.value == 'RTA Address':
            rta_addr_col_idx = cell.column
        elif cell.value == 'RTA Status':
            rta_status_col_idx = cell.column

    yellow_fill = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
    red_fill = PatternFill(start_color='FF6666', end_color='FF6666', fill_type='solid')
    orange_fill = PatternFill(start_color='FFA500', end_color='FFA500', fill_type='solid')

    for i, mt in enumerate(match_type_series):
        if mt in ('fuzzy', 'direction_strip'):
            fill = yellow_fill
        elif mt == 'no_pc':
            fill = red_fill
        elif mt == 'conflict':
            fill = orange_fill
        else:
            continue
        if rta_addr_col_idx:
            ws_hub.cell(row=i + 2, column=rta_addr_col_idx).fill = fill
        if rta_status_col_idx:
            ws_hub.cell(row=i + 2, column=rta_status_col_idx).fill = fill

    # Color RTA "Not in Hubspot" rows purple
    ws_rta = wb['RTA']
    purple_fill = PatternFill(start_color='D8B4FE', end_color='D8B4FE', fill_type='solid')
    in_hub_col_idx = None
    for cell in ws_rta[1]:
        if cell.value == 'In Hubspot':
            in_hub_col_idx = cell.column
            break
    if in_hub_col_idx:
        for row_idx in range(2, ws_rta.max_row + 1):
            cell = ws_rta.cell(row=row_idx, column=in_hub_col_idx)
            if cell.value == 'No':
                for col_idx in range(1, ws_rta.max_column + 1):
                    ws_rta.cell(row=row_idx, column=col_idx).fill = purple_fill

    out_buffer = io.BytesIO()
    wb.save(out_buffer)
    excel_bytes = out_buffer.getvalue()

    return {
        'stats': {
            'hubspot_total': len(df_hub),
            'hubspot_matched': hub_matched,
            'hubspot_unmatched': hub_unmatched,
            'rta_total': len(df_rta),
            'rta_in_hubspot': rta_in_hub,
            'rta_not_in_hubspot': rta_not_in_hub,
            'exact': exact_count,
            'fuzzy': yellow_count,
            'conflict': orange_count,
            'risky_no_pc': red_count,
        },
        'conflicts': conflict_detail,
        'flagged': flagged,
        'rta_not_in_hubspot_preview': rta_not_in_hubspot_preview,
        'hub_output_preview': hub_output_preview,
        'excel_bytes': excel_bytes,
        # DataFrames with _k_* normalization columns intact — caller pickles for
        # later search via app.services.search.search_in_dataset.
        'df_hub_keyed': df_hub,
        'df_rta_keyed': df_rta,
    }
